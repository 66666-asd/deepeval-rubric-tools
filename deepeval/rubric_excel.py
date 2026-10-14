"""Downstream Excel output for validated, already-scored rubric comparisons."""

from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path

from deepeval.rubric import RubricValidationError

_INVALID_XML_TEXT = re.compile(
    r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]"
)
_MAX_EXCEL_ROWS = 1048576


def _new_workbook():
    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise RubricValidationError(
            "Excel output is optional. Install requirements-rubric-excel.txt."
        ) from exc
    return Workbook()


def _append_literals(sheet, values):
    """Write user strings as literal text, never executable Excel formulas."""
    if sheet.max_row >= _MAX_EXCEL_ROWS:
        raise RubricValidationError("Report exceeds the Excel row limit.")
    for value in values:
        if isinstance(value, str):
            if _INVALID_XML_TEXT.search(value):
                raise RubricValidationError(
                    "Report text contains characters not supported by Excel."
                )
            if len(value.encode("utf-16-le")) // 2 > 32767:
                raise RubricValidationError(
                    "Report text exceeds the Excel cell limit (32767 units)."
                )
    sheet.append(values)
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = "s"


def _style_sheet(sheet, widths, numeric_columns=(), filtered=True):
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    edge = Side(style="thin", color="D0D5DD")
    for cell in sheet[1]:
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1D2330")
        cell.border = Border(bottom=edge)
    sheet.row_dimensions[1].height = 24
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Calibri", size=11, color="222222")
            cell.alignment = Alignment(
                vertical="top",
                horizontal="right"
                if cell.column in numeric_columns
                else "left",
            )
            if cell.column in numeric_columns and isinstance(
                cell.value, (int, float)
            ):
                cell.number_format = "0.00"
    if filtered:
        sheet.auto_filter.ref = sheet.dimensions


def _save_new_workbook(workbook, output):
    """Serialize first; exclusive create never overwrites an existing file."""
    output = Path(output)
    if output.suffix.lower() != ".xlsx":
        raise RubricValidationError("Excel output path must end in .xlsx.")
    buffer = BytesIO()
    try:
        workbook.save(buffer)
        with output.open("xb") as target:
            target.write(buffer.getbuffer())
    finally:
        buffer.close()
        workbook.close()


def write_rubric_excel_report(report, output):
    """Export a snapshot of validated scores without changing their precision.

    These are result values, not spreadsheet calculations. The input data and
    JSON calculation retain full precision; two decimals are display-only.
    """
    workbook = _new_workbook()
    try:
        summary = workbook.active
        summary.title = "Summary"
        _append_literals(summary, ["Metric", "Value"])
        summary_rows = [
            ("Model A", report["models"]["a"]),
            ("Model B", report["models"]["b"]),
            ("Model A score (0-100)", report["model_a_score"]),
            ("Model B score (0-100)", report["model_b_score"]),
            (
                "A minus B (percentage points)",
                report["delta_percentage_points"],
            ),
            ("Winner", report["winner"]),
            ("Cases", report["case_count"]),
            ("Criteria", report["criterion_count"]),
            ("Method", report["method"]),
            ("Source", report["source"]),
            ("Source worksheet", report.get("source_sheet", "Not specified")),
            ("Schema version", report["schema_version"]),
            (
                "Tie tolerance (percentage points)",
                report["tie_tolerance_percentage_points"],
            ),
        ]
        manifest = report.get("manifest_check")
        summary_rows.append(
            (
                "Manifest check",
                "Not requested" if manifest is None else manifest["status"],
            )
        )
        if manifest is not None:
            summary_rows.extend(
                [
                    ("Manifest source", manifest["source"]),
                    ("Expected criteria", manifest["expected_criterion_count"]),
                    ("Matched criteria", manifest["matched_criterion_count"]),
                ]
            )
        for row in summary_rows:
            _append_literals(summary, row)
        _style_sheet(summary, [38, 100], filtered=False)
        for row in (4, 5, 6):
            summary.cell(row, 2).number_format = "0.00"

        cases = workbook.create_sheet("Cases")
        _append_literals(
            cases,
            [
                "case_id",
                "criterion_count",
                "model_a_score",
                "model_b_score",
                "delta_percentage_points",
                "winner",
            ],
        )
        criteria = workbook.create_sheet("Criteria")
        fields = (
            "criterion_id",
            "criterion",
            "weight",
            "weight_share",
            "max_score",
            "model_a_score",
            "model_b_score",
            "model_a_contribution",
            "model_b_contribution",
            "delta_percentage_points",
            "winner",
            "notes",
        )
        _append_literals(criteria, ["case_id", *fields])
        for case in report["cases"]:
            _append_literals(
                cases,
                [
                    case["case_id"],
                    case["criterion_count"],
                    case["model_a_score"],
                    case["model_b_score"],
                    case["delta_percentage_points"],
                    case["winner"],
                ],
            )
            for criterion in case["criteria"]:
                _append_literals(
                    criteria,
                    [case["case_id"], *(criterion[field] for field in fields)],
                )
        _style_sheet(cases, [22, 18, 22, 22, 28, 26], numeric_columns=(3, 4, 5))
        _style_sheet(
            criteria,
            [22, 22, 55, 16, 18, 16, 20, 20, 24, 24, 28, 26, 65],
            numeric_columns=range(4, 12),
        )
        for row in criteria.iter_rows(min_row=2):
            row[4].number_format = "0.0%"
        _save_new_workbook(workbook, output)
    finally:
        workbook.close()
