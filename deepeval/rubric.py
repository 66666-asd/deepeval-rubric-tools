"""Validate and compare existing human rubric scores without a model call."""

from __future__ import annotations

import csv
import math
from collections.abc import Iterable
from dataclasses import dataclass
from pathlib import Path
from typing import Any
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

REQUIRED_COLUMNS = (
    "case_id",
    "criterion_id",
    "criterion",
    "weight",
    "max_score",
    "model_a_score",
    "model_b_score",
)
OPTIONAL_COLUMNS = ("notes",)
TEMPLATE_REQUIRED_COLUMNS = (
    "case_id",
    "criterion_id",
    "criterion",
    "weight",
    "max_score",
)
TEMPLATE_OPTIONAL_COLUMNS = ("notes",)
TIE_TOLERANCE = 1e-9


class RubricValidationError(ValueError):
    """A score sheet cannot be safely compared."""


@dataclass(frozen=True)
class RubricScore:
    case_id: str
    criterion_id: str
    criterion: str
    weight: float
    max_score: float
    model_a_score: float
    model_b_score: float
    notes: str = ""


def _number(value: str, column: str, line: int) -> float:
    try:
        result = float(value)
    except ValueError:
        raise RubricValidationError(
            f"Line {line}: {column} must be a number; got {value!r}."
        ) from None
    if not math.isfinite(result):
        raise RubricValidationError(f"Line {line}: {column} must be finite.")
    return result


def _validate_columns(columns, required, optional=()):
    if not columns:
        raise RubricValidationError("Table is empty or has no header.")
    if not all(isinstance(column, str) for column in columns):
        raise RubricValidationError("Headers must be nonempty text.")
    if len(columns) != len(set(columns)):
        raise RubricValidationError("Table contains duplicate headers.")
    missing = sorted(set(required) - set(columns))
    unknown = sorted(set(columns) - set(required) - set(optional))
    if missing or unknown:
        raise RubricValidationError(
            f"Invalid columns: missing={missing}; unknown={unknown}."
        )


def _read_csv(path, required=REQUIRED_COLUMNS, optional=OPTIONAL_COLUMNS):
    try:
        with Path(path).open(encoding="utf-8-sig", newline="") as source:
            reader = csv.DictReader(source, strict=True)
            _validate_columns(reader.fieldnames, required, optional)
            for record in reader:
                line = reader.line_num
                if None in record or any(
                    value is None for value in record.values()
                ):
                    raise RubricValidationError(
                        f"Line {line}: field count does not match the header."
                    )
                yield (
                    line,
                    {key: value.strip() for key, value in record.items()},
                )
    except (UnicodeError, csv.Error) as exc:
        raise RubricValidationError(
            f"Cannot read CSV as valid UTF-8 CSV: {exc}"
        ) from exc


def _score_records(records: Iterable[tuple[int, dict[str, str]]]):
    rows: list[RubricScore] = []
    seen = set()
    for line, record in records:
        for column in REQUIRED_COLUMNS:
            if not record[column]:
                raise RubricValidationError(
                    f"Line {line}: {column} is required."
                )
        key = (record["case_id"], record["criterion_id"])
        if key in seen:
            raise RubricValidationError(
                f"Line {line}: duplicate case/criterion {key!r}."
            )
        seen.add(key)
        numbers = {
            column: _number(record[column], column, line)
            for column in REQUIRED_COLUMNS[3:]
        }
        for column in ("weight", "max_score"):
            if numbers[column] <= 0:
                raise RubricValidationError(
                    f"Line {line}: {column} must be greater than zero."
                )
        for column in ("model_a_score", "model_b_score"):
            if not 0 <= numbers[column] <= numbers["max_score"]:
                raise RubricValidationError(
                    f"Line {line}: {column} must be between zero and max_score."
                )
        rows.append(
            RubricScore(
                case_id=key[0],
                criterion_id=key[1],
                criterion=record["criterion"],
                notes=record.get("notes", ""),
                **numbers,
            )
        )
    if not rows:
        raise RubricValidationError(
            "Table must contain at least one score row."
        )
    return rows


def load_rubric_csv(path: str | Path) -> list[RubricScore]:
    """Load UTF-8 (optionally BOM) CSV; never turn missing scores into zero."""
    return _score_records(_read_csv(path))


def load_rubric_template_csv(path: str | Path) -> list[dict[str, Any]]:
    """Load Rubric definitions for a blank, import-compatible score template."""
    if Path(path).suffix.lower() != ".csv":
        raise RubricValidationError(
            "Rubric definition must be a UTF-8 CSV file."
        )
    rows = []
    seen = set()
    for line, record in _read_csv(
        path,
        required=TEMPLATE_REQUIRED_COLUMNS,
        optional=TEMPLATE_OPTIONAL_COLUMNS,
    ):
        if len(rows) >= 100000:
            raise RubricValidationError(
                "Rubric definition exceeds 100000 data rows."
            )
        for column in TEMPLATE_REQUIRED_COLUMNS[:3]:
            if not record[column]:
                raise RubricValidationError(
                    f"Line {line}: {column} is required."
                )
        key = (record["case_id"], record["criterion_id"])
        if key in seen:
            raise RubricValidationError(
                f"Line {line}: duplicate case/criterion {key!r}."
            )
        seen.add(key)
        weight = _number(record["weight"], "weight", line)
        max_score = _number(record["max_score"], "max_score", line)
        if weight <= 0 or max_score <= 0:
            raise RubricValidationError(
                f"Line {line}: weight and max_score must be greater than zero."
            )
        rows.append(
            {
                "case_id": key[0],
                "criterion_id": key[1],
                "criterion": record["criterion"],
                "weight": weight,
                "max_score": max_score,
                "notes": record.get("notes", ""),
            }
        )
    if not rows:
        raise RubricValidationError(
            "Rubric definition must contain at least one row."
        )
    return rows


def _read_xlsx(path, sheet):
    try:
        import defusedxml  # noqa: F401 - required for safe XML parsing.
        from defusedxml.common import DefusedXmlException
        from openpyxl import load_workbook
        from openpyxl.utils.exceptions import InvalidFileException
        from openpyxl.xml import DEFUSEDXML
    except ImportError as exc:
        raise RubricValidationError(
            "Excel support is optional. Install requirements-rubric-excel.txt."
        ) from exc
    if not DEFUSEDXML:
        raise RubricValidationError("Excel input requires defusedxml enabled.")
    path = Path(path)
    if path.stat().st_size > 50 * 1024 * 1024:
        raise RubricValidationError("Excel file exceeds the 50 MiB size limit.")
    try:
        with ZipFile(path) as archive:
            entries = archive.infolist()
            if (
                len(entries) > 10000
                or sum(e.file_size for e in entries) > 200 * 1024 * 1024
            ):
                raise RubricValidationError(
                    "Excel archive exceeds safety limits."
                )
        workbook = load_workbook(
            path,
            read_only=True,
            data_only=False,
            keep_links=False,
        )
        try:
            if sheet is None:
                if len(workbook.sheetnames) != 1:
                    raise RubricValidationError(
                        "Workbook has multiple sheets; select one with --sheet."
                    )
                sheet = workbook.sheetnames[0]
            if sheet not in workbook.sheetnames:
                raise RubricValidationError(
                    f"Worksheet {sheet!r} was not found."
                )
            worksheet = workbook[sheet]
            # Read actual cells instead of trusting possibly stale dimensions.
            worksheet.reset_dimensions()
            iterator = worksheet.iter_rows()
            header = next(iterator, ())
            columns = [cell.value for cell in header]
            _validate_columns(columns, REQUIRED_COLUMNS, OPTIONAL_COLUMNS)
            for line, cells in enumerate(iterator, start=2):
                if line > 100001:
                    raise RubricValidationError(
                        "Excel sheet exceeds 100000 data rows."
                    )
                if all(cell.value is None for cell in cells):
                    continue
                if any(
                    cell.value is not None for cell in cells[len(columns) :]
                ):
                    raise RubricValidationError(
                        f"Line {line}: unexpected extra cells."
                    )
                record = {}
                for index, column in enumerate(columns):
                    cell = cells[index] if index < len(cells) else None
                    value = cell.value if cell is not None else None
                    if cell is not None and cell.data_type in ("f", "e"):
                        raise RubricValidationError(
                            f"Cell {cell.coordinate}: formulas and Excel errors are not allowed."
                        )
                    if isinstance(value, bool):
                        raise RubricValidationError(
                            f"Line {line}: booleans are not scores or identifiers."
                        )
                    if (
                        value is not None
                        and column in REQUIRED_COLUMNS[:3] + OPTIONAL_COLUMNS
                        and not isinstance(value, str)
                    ):
                        raise RubricValidationError(
                            f"Line {line}: {column} must be text (preserve leading zeros)."
                        )
                    record[column] = "" if value is None else str(value).strip()
                yield line, record
        finally:
            workbook.close()
    except (
        BadZipFile,
        ParseError,
        InvalidFileException,
        DefusedXmlException,
        KeyError,
        ValueError,
    ) as exc:
        if isinstance(exc, RubricValidationError):
            raise
        raise RubricValidationError(
            f"Cannot read a valid XLSX workbook: {exc}"
        ) from exc


def load_rubric_scores(path: str | Path, sheet: str | None = None):
    """Read a CSV or XLSX sheet using identical score validation."""
    suffix = Path(path).suffix.lower()
    if suffix == ".xlsx":
        return _score_records(_read_xlsx(path, sheet))
    if suffix == ".csv":
        if sheet is not None:
            raise RubricValidationError(
                "--sheet can only be used with XLSX input."
            )
        return load_rubric_csv(path)
    raise RubricValidationError("Supported input formats are .csv and .xlsx.")


def _winner(delta: float, model_a: str, model_b: str) -> str:
    if abs(delta) <= TIE_TOLERANCE:
        return "tie"
    return model_a if delta > 0 else model_b


def compare_rubric_csv(
    path: str | Path,
    model_a: str = "model_a",
    model_b: str = "model_b",
) -> dict[str, Any]:
    """Backward-compatible CSV comparison API."""
    return _compare_scores(load_rubric_csv(path), path, model_a, model_b)


def compare_rubric_file(
    path: str | Path,
    model_a: str = "model_a",
    model_b: str = "model_b",
    *,
    sheet: str | None = None,
    manifest: str | Path | None = None,
) -> dict[str, Any]:
    """Compare CSV or XLSX scores without changing the scoring convention."""
    rows = load_rubric_scores(path, sheet)
    manifest_check = None
    if manifest is not None:
        manifest_check = validate_rubric_manifest(rows, manifest)
    report = _compare_scores(rows, path, model_a, model_b)
    if sheet is not None:
        report["source_sheet"] = sheet
    if manifest_check is not None:
        report["manifest_check"] = manifest_check
    return report


def validate_rubric_manifest(
    rows: list[RubricScore], path: str | Path
) -> dict[str, Any]:
    """Require an exact match to an independently supplied CSV key manifest.

    No scores or scoring policy are inferred from the manifest. A manifest
    constructed from the score file cannot detect jointly omitted criteria.
    """
    if Path(path).suffix.lower() != ".csv":
        raise RubricValidationError("Manifest must be a UTF-8 CSV file.")
    columns = ("case_id", "criterion_id")
    expected = set()
    for line, record in _read_csv(path, required=columns, optional=()):
        for column in columns:
            if not record[column]:
                raise RubricValidationError(
                    f"Manifest line {line}: {column} is required."
                )
        key = (record["case_id"], record["criterion_id"])
        if key in expected:
            raise RubricValidationError(
                f"Manifest line {line}: duplicate case/criterion {key!r}."
            )
        expected.add(key)
    if not expected:
        raise RubricValidationError("Manifest must contain at least one key.")
    actual = {(row.case_id, row.criterion_id) for row in rows}
    missing = sorted(expected - actual)
    unexpected = sorted(actual - expected)
    if missing or unexpected:
        raise RubricValidationError(
            "Manifest mismatch (first 20 keys shown per group): "
            f"missing={len(missing)} {missing[:20]!r}; "
            f"unexpected={len(unexpected)} {unexpected[:20]!r}."
        )
    return {
        "source": str(Path(path)),
        "status": "passed",
        "expected_criterion_count": len(expected),
        "matched_criterion_count": len(actual),
    }


def _compare_scores(rows, path, model_a, model_b):
    """Compute per-case weighted percentages, then an equal-case macro mean.

    Input scores are supplied judgments, not judgments made by DeepEval.
    A positive delta always means model A has the higher score.
    """
    model_a, model_b = model_a.strip(), model_b.strip()
    if (
        not model_a
        or not model_b
        or model_a == model_b
        or "tie" in (model_a, model_b)
    ):
        raise RubricValidationError(
            "Model names must be nonempty, distinct, and not 'tie'."
        )
    groups: dict[str, list[RubricScore]] = {}
    for row in rows:
        groups.setdefault(row.case_id, []).append(row)

    cases = []
    wins = {model_a: 0, model_b: 0, "tie": 0}
    for case_id, criteria in groups.items():
        # Scale before summing so large, finite weights cannot overflow.
        scale = max(row.weight for row in criteria)
        total_weight = math.fsum(row.weight / scale for row in criteria)
        details = []
        for row in criteria:
            share = (row.weight / scale) / total_weight
            a = row.model_a_score / row.max_score * 100
            b = row.model_b_score / row.max_score * 100
            details.append(
                {
                    "criterion_id": row.criterion_id,
                    "criterion": row.criterion,
                    "weight": row.weight,
                    "weight_share": share,
                    "max_score": row.max_score,
                    "model_a_score": row.model_a_score,
                    "model_b_score": row.model_b_score,
                    "model_a_contribution": share * a,
                    "model_b_contribution": share * b,
                    "delta_percentage_points": a - b,
                    "winner": _winner(a - b, model_a, model_b),
                    "notes": row.notes,
                }
            )
        a = math.fsum(item["model_a_contribution"] for item in details)
        b = math.fsum(item["model_b_contribution"] for item in details)
        winner = _winner(a - b, model_a, model_b)
        wins[winner] += 1
        cases.append(
            {
                "case_id": case_id,
                "criterion_count": len(criteria),
                "model_a_score": a,
                "model_b_score": b,
                "delta_percentage_points": a - b,
                "winner": winner,
                "criteria": details,
            }
        )
    a = math.fsum(case["model_a_score"] for case in cases) / len(cases)
    b = math.fsum(case["model_b_score"] for case in cases) / len(cases)
    return {
        "schema_version": 1,
        "source": str(Path(path)),
        "models": {"a": model_a, "b": model_b},
        "method": "weighted normalized criteria; equal-weight case macro mean",
        "score_range": [0, 100],
        "tie_tolerance_percentage_points": TIE_TOLERANCE,
        "case_count": len(cases),
        "criterion_count": len(rows),
        "model_a_score": a,
        "model_b_score": b,
        "delta_percentage_points": a - b,
        "winner": _winner(a - b, model_a, model_b),
        "case_wins": wins,
        "cases": cases,
    }
