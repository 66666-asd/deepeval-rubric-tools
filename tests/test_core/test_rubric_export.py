"""Excel exports preserve results, literal text, and existing user files."""

import csv
import json

import pytest
from openpyxl import load_workbook
from typer.testing import CliRunner

from deepeval.cli.main import app
from deepeval.rubric import (
    REQUIRED_COLUMNS,
    RubricValidationError,
    compare_rubric_file,
)
from deepeval.rubric_excel import write_rubric_excel_report


@pytest.fixture()
def source(tmp_path):
    path = tmp_path / "scores.csv"
    with path.open("w", newline="", encoding="utf-8") as target:
        writer = csv.writer(target)
        writer.writerow([*REQUIRED_COLUMNS, "notes"])
        writer.writerows(
            [
                ["001", "r1", "Accuracy", 3, 4, 4, 2, "Synthetic scores"],
                ["001", "r2", "Evidence", 1, 2, 1, 2, ""],
                ["002", "r1", "Coverage", 1, 3, 1, 2, ""],
            ]
        )
    return path


def test_export_matches_report_and_keeps_full_precision(
    tmp_path, source, socket_disabled
):
    report = compare_rubric_file(source, "Baseline", "Candidate")
    output = tmp_path / "report.xlsx"
    write_rubric_excel_report(report, output)
    book = load_workbook(output, data_only=False)
    try:
        assert book.sheetnames == ["Summary", "Cases", "Criteria"]
        summary = dict(book["Summary"].iter_rows(min_row=2, values_only=True))
        assert summary["Model A score (0-100)"] == pytest.approx(
            report["model_a_score"]
        )
        assert summary["Model A"] == "Baseline"
        assert summary["Manifest check"] == "Not requested"
        assert book["Cases"].max_row == 3
        assert book["Criteria"].max_row == 4
        assert book["Cases"]["A2"].value == "001"
        assert book["Cases"]["A2"].data_type == "s"
        assert book["Cases"]["C3"].value == pytest.approx(100 / 3)
        assert book["Cases"]["C3"].number_format == "0.00"
        assert book["Criteria"]["E2"].value == 0.75
        assert book["Criteria"]["E2"].number_format == "0.0%"
        assert book["Cases"].auto_filter.ref == "A1:F3"
        assert all(
            cell.data_type != "f"
            for sheet in book
            for row in sheet
            for cell in row
        )
    finally:
        book.close()


@pytest.mark.parametrize(
    "text",
    [
        '=HYPERLINK("https://example.invalid","x")',
        "+1+1",
        "-1+1",
        "@SUM(A1)",
        "#REF!",
        "\u4e2d\u6587\nreview",
    ],
)
def test_user_strings_are_literals_not_formulas(tmp_path, source, text):
    report = compare_rubric_file(source, model_a=text)
    report["cases"][0]["criteria"][0]["notes"] = text
    output = tmp_path / "report.xlsx"
    write_rubric_excel_report(report, output)
    book = load_workbook(output)
    try:
        assert book["Summary"]["B2"].value == text
        assert book["Summary"]["B2"].data_type == "s"
        assert book["Criteria"]["M2"].value == text
        assert book["Criteria"]["M2"].data_type == "s"
        assert book["Criteria"]["M2"].hyperlink is None
    finally:
        book.close()


@pytest.mark.parametrize(
    "text",
    ["x" * 32768, "bad\x00text", "bad\ud800text", "\uffff"],
    ids=["too_long", "nul", "surrogate", "noncharacter"],
)
def test_invalid_excel_text_leaves_no_partial_output(tmp_path, source, text):
    report = compare_rubric_file(source)
    report["cases"][0]["criteria"][0]["notes"] = text
    output = tmp_path / "report.xlsx"
    with pytest.raises(RubricValidationError):
        write_rubric_excel_report(report, output)
    assert not output.exists()


def test_existing_file_is_unchanged(tmp_path, source):
    output = tmp_path / "report.xlsx"
    output.write_bytes(b"existing user content")
    with pytest.raises(FileExistsError):
        write_rubric_excel_report(compare_rubric_file(source), output)
    assert output.read_bytes() == b"existing user content"


def test_cli_exports_excel_and_original_json_still_works(
    tmp_path, source, socket_disabled
):
    output = tmp_path / "report.xlsx"
    runner = CliRunner()
    result = runner.invoke(
        app, ["rubric", "compare", str(source), "-o", str(output)]
    )
    assert result.exit_code == 0, result.output
    assert output.read_bytes().startswith(b"PK")
    result = runner.invoke(app, ["rubric", "compare", str(source)])
    assert result.exit_code == 0
    assert json.loads(result.stdout)["case_count"] == 2


def test_cli_rejects_unknown_export_extension(tmp_path, source):
    output = tmp_path / "report.xls"
    result = CliRunner().invoke(
        app, ["rubric", "compare", str(source), "-o", str(output)]
    )
    assert result.exit_code == 2
    assert not output.exists()


def test_manifest_details_are_exported(tmp_path, source):
    manifest = tmp_path / "manifest.csv"
    manifest.write_text(
        "case_id,criterion_id\n001,r1\n001,r2\n002,r1\n", encoding="utf-8"
    )
    report = compare_rubric_file(source, manifest=manifest)
    output = tmp_path / "report.xlsx"
    write_rubric_excel_report(report, output)
    book = load_workbook(output)
    try:
        summary = dict(book["Summary"].iter_rows(min_row=2, values_only=True))
        assert summary["Manifest check"] == "passed"
        assert summary["Expected criteria"] == 3
        assert summary["Matched criteria"] == 3
    finally:
        book.close()
