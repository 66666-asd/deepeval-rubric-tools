"""Definition-driven templates remain blank until human scores are supplied."""

import csv
from pathlib import Path

import pytest
from openpyxl import load_workbook
from typer.testing import CliRunner

from deepeval.cli.main import app
from deepeval.rubric import (
    REQUIRED_COLUMNS,
    TEMPLATE_REQUIRED_COLUMNS,
    RubricValidationError,
    compare_rubric_file,
    load_rubric_template_csv,
)
from deepeval.rubric_excel import write_rubric_excel_template


def write_definition(path, rows, headers=(*TEMPLATE_REQUIRED_COLUMNS, "notes")):
    with path.open("w", encoding="utf-8", newline="") as target:
        writer = csv.writer(target)
        writer.writerow(headers)
        writer.writerows(rows)
    return path


@pytest.fixture()
def definition(tmp_path):
    return write_definition(
        tmp_path / "definition.csv",
        [
            ["001", "r1", "Accuracy", 3, 4, "Synthetic definition"],
            ["001", "r2", "Evidence", 1, 2, ""],
            ["002", "r1", "Coverage", 1, 3, ""],
        ],
    )


def test_template_has_blank_scores_text_ids_and_range_validation(
    definition, tmp_path, socket_disabled
):
    output = tmp_path / "blank.xlsx"
    assert write_rubric_excel_template(definition, output) == 3
    book = load_workbook(output)
    try:
        assert book.sheetnames == ["Scores", "Instructions"]
        scores = book["Scores"]
        assert [cell.value for cell in scores[1]] == [
            *REQUIRED_COLUMNS,
            "notes",
        ]
        assert scores.max_row == 4
        assert scores["A2"].value == "001"
        assert scores["A2"].data_type == "s"
        assert scores["A2"].number_format == "@"
        assert scores["D2"].value == 3
        assert scores["E2"].value == 4
        assert scores["H2"].value == "Synthetic definition"
        for row in scores.iter_rows(min_row=2):
            assert row[5].value is None and row[6].value is None
        assert scores.auto_filter.ref == "A1:H4"
        validation = scores.data_validations.dataValidation[0]
        assert str(validation.sqref) == "F2:G4"
        assert validation.type == "decimal"
        assert validation.operator == "between"
        assert validation.formula1 == "0"
        assert validation.formula2 == "$E2"
        assert validation.allowBlank and validation.showErrorMessage
        assert validation.errorStyle == "stop"
        assert all(
            cell.data_type != "f"
            for sheet in book
            for row in sheet
            for cell in row
        )
    finally:
        book.close()


def test_blank_template_rejected_then_completed_template_imports(
    definition, tmp_path
):
    output = tmp_path / "scores.xlsx"
    write_rubric_excel_template(definition, output)
    with pytest.raises(
        RubricValidationError, match="model_a_score is required"
    ):
        compare_rubric_file(output, sheet="Scores")
    book = load_workbook(output)
    try:
        for row in book["Scores"].iter_rows(min_row=2):
            row[5].value = row[4].value
            row[6].value = 0
        book.save(output)
    finally:
        book.close()
    manifest = tmp_path / "manifest.csv"
    manifest.write_text(
        "case_id,criterion_id\n001,r1\n001,r2\n002,r1\n", encoding="utf-8"
    )
    report = compare_rubric_file(output, sheet="Scores", manifest=manifest)
    assert report["model_a_score"] == 100
    assert report["model_b_score"] == 0
    assert report["criterion_count"] == 3
    assert report["manifest_check"]["status"] == "passed"
    with pytest.raises(RubricValidationError, match="multiple sheets"):
        compare_rubric_file(output)


@pytest.mark.parametrize("score", [-1, 5, "=1+1", True])
def test_import_rechecks_values_pasted_over_excel_validation(
    definition, tmp_path, score
):
    output = tmp_path / "scores.xlsx"
    write_rubric_excel_template(definition, output)
    book = load_workbook(output)
    try:
        for row in book["Scores"].iter_rows(min_row=2):
            row[5].value = row[6].value = 0
        book["Scores"]["F2"].value = score
        book.save(output)
    finally:
        book.close()
    with pytest.raises(RubricValidationError):
        compare_rubric_file(output, sheet="Scores")


@pytest.mark.parametrize(
    "field,value",
    [
        ("case_id", ""),
        ("criterion_id", ""),
        ("criterion", ""),
        ("weight", ""),
        ("weight", "nan"),
        ("weight", "inf"),
        ("weight", 0),
        ("weight", -1),
        ("max_score", 0),
        ("max_score", -1),
        ("max_score", "NaN"),
        ("max_score", "text"),
    ],
)
def test_invalid_definition_prevents_file_creation(tmp_path, field, value):
    row = ["001", "r1", "Accuracy", 1, 4]
    row[TEMPLATE_REQUIRED_COLUMNS.index(field)] = value
    definition = write_definition(
        tmp_path / "bad.csv", [row], TEMPLATE_REQUIRED_COLUMNS
    )
    output = tmp_path / "scores.xlsx"
    with pytest.raises(RubricValidationError):
        write_rubric_excel_template(definition, output)
    assert not output.exists()


@pytest.mark.parametrize(
    "content",
    [
        "",
        "case_id,criterion_id,criterion,weight,max_score\n",
        "case_id,criterion_id,criterion,weight,max_score\na,r,c,1,4\na,r,c,1,4\n",
        "case_id,criterion_id,criterion,weight,max_score\na,r,c,1\n",
        "case_id,criterion_id,criterion,weight,max_score\na,r,c,1,4,extra\n",
        "case_id,criterion_id,criterion,weight,weight\na,r,c,1,4\n",
        "case_id,criterion_id,criterion,weight,max_score,model_a_score\na,r,c,1,4,2\n",
    ],
    ids=[
        "empty",
        "no_rows",
        "duplicate",
        "short",
        "long",
        "duplicate_header",
        "scores",
    ],
)
def test_invalid_csv_definitions_are_rejected(tmp_path, content):
    definition = tmp_path / "bad.csv"
    definition.write_text(content, encoding="utf-8")
    with pytest.raises(RubricValidationError):
        load_rubric_template_csv(definition)


def test_definition_accepts_bom_reordered_headers_and_optional_notes(tmp_path):
    definition = tmp_path / "definition.CSV"
    definition.write_text(
        "max_score,criterion_id,weight,criterion,case_id\n4,r1,2,\u4e2d\u6587,001\n",
        encoding="utf-8-sig",
    )
    assert load_rubric_template_csv(definition) == [
        {
            "case_id": "001",
            "criterion_id": "r1",
            "criterion": "\u4e2d\u6587",
            "weight": 2.0,
            "max_score": 4.0,
            "notes": "",
        }
    ]


def test_formula_like_definition_text_remains_literal(tmp_path):
    text = '=HYPERLINK("https://example.invalid", "x")'
    definition = write_definition(
        tmp_path / "definition.csv", [["001", "r1", text, 1, 4, text]]
    )
    output = tmp_path / "blank.xlsx"
    write_rubric_excel_template(definition, output)
    book = load_workbook(output)
    try:
        for coordinate in ("C2", "H2"):
            assert book["Scores"][coordinate].value == text
            assert book["Scores"][coordinate].data_type == "s"
    finally:
        book.close()


def test_existing_output_and_definition_are_not_overwritten(
    definition, tmp_path
):
    original = definition.read_bytes()
    output = tmp_path / "existing.xlsx"
    output.write_bytes(b"existing user content")
    with pytest.raises(FileExistsError):
        write_rubric_excel_template(definition, output)
    assert output.read_bytes() == b"existing user content"
    with pytest.raises(RubricValidationError, match=".xlsx"):
        write_rubric_excel_template(definition, definition)
    assert definition.read_bytes() == original


def test_cli_creates_template_and_reports_file_errors(
    definition, tmp_path, socket_disabled
):
    output = tmp_path / "blank.xlsx"
    runner = CliRunner()
    args = ["rubric", "template", str(definition), "-o", str(output)]
    result = runner.invoke(app, args)
    assert result.exit_code == 0, result.output
    assert "3 criteria" in result.stdout
    original = output.read_bytes()
    result = runner.invoke(app, args)
    assert result.exit_code == 2
    assert output.read_bytes() == original
    result = runner.invoke(app, ["rubric", "template", str(definition)])
    assert result.exit_code == 2


@pytest.mark.parametrize("name", ["definition.xlsx", "absent.csv"])
def test_cli_rejects_wrong_or_missing_definition(tmp_path, name):
    output = tmp_path / "blank.xlsx"
    result = CliRunner().invoke(
        app, ["rubric", "template", str(tmp_path / name), "-o", str(output)]
    )
    assert result.exit_code == 2
    assert not output.exists()


@pytest.mark.parametrize(
    "text", ["bad\x00text", "x" * 32768], ids=["nul", "too_long"]
)
def test_invalid_excel_definition_text_creates_no_output(tmp_path, text):
    definition = write_definition(
        tmp_path / "definition.csv", [["001", "r1", text, 1, 4, ""]]
    )
    output = tmp_path / "blank.xlsx"
    with pytest.raises(RubricValidationError):
        write_rubric_excel_template(definition, output)
    assert not output.exists()


def test_definition_row_limit_is_enforced(monkeypatch):
    record = {
        "case_id": "a",
        "criterion": "Accuracy",
        "weight": "1",
        "max_score": "4",
    }
    monkeypatch.setattr(
        "deepeval.rubric._read_csv",
        lambda *args, **kwargs: (
            (index + 2, {**record, "criterion_id": str(index)})
            for index in range(100001)
        ),
    )
    with pytest.raises(RubricValidationError, match="100000"):
        load_rubric_template_csv("definition.csv")


def test_documented_example_definitions_match_independent_manifest(tmp_path):
    examples = Path(__file__).resolve().parents[2] / "examples" / "rubric"
    output = tmp_path / "example.xlsx"
    assert (
        write_rubric_excel_template(examples / "definitions.csv", output) == 4
    )
    book = load_workbook(output)
    try:
        for row in book["Scores"].iter_rows(min_row=2):
            row[5].value = row[6].value = 0
        book.save(output)
    finally:
        book.close()
    report = compare_rubric_file(
        output, sheet="Scores", manifest=examples / "expected_criteria.csv"
    )
    assert report["manifest_check"]["matched_criterion_count"] == 4
